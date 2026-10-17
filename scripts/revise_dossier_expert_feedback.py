from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "AI2026 - Hồ sơ Bảng A.docx"
OUTPUT = ROOT / "AI2026 - Hồ sơ Bảng A - Bản chỉnh sửa theo phản biện.docx"


STATUS_ROWS = [
    (
        "Phạm vi và số liệu đường cơ sở",
        "Ghi rõ sản phẩm hiện dùng kho tri thức Lịch sử lớp 8; giữ nguyên kết quả "
        "kiểm thử hồi quy 59/60 câu đạt và thời gian trung bình 23,32 giây.",
        "Tách kết quả đã đo khỏi cam kết tương lai; tránh hiểu nhầm sản phẩm đã "
        "bao phủ các khối lớp khác.",
        "Đã hoàn thành",
    ),
    (
        "Nguồn và giới hạn dữ liệu",
        "Giữ thông tin nguồn và trạng thái thẩm định cùng bản ghi; chỉ hiển thị số trang "
        "SGK khi metadata đã được xác minh, nếu chưa có thì nói rõ chưa ghi nhận.",
        "Người học biết căn cứ nào đã có và phần nào chưa được kiểm chứng; "
        "không bịa nguồn hoặc số trang.",
        "Đã hoàn thành",
    ),
    (
        "Đoạn dữ liệu dùng để đối chiếu cạnh câu trả lời",
        "Danh sách nguồn hiển thị nội dung của bản ghi đã truy xuất, tên nguồn và trạng thái thẩm "
        "định; giao diện nói rõ nội dung HISEDU không phải trích dẫn nguyên văn SGK khi trang chưa xác minh.",
        "Giúp phát hiện nhanh lỗi sai niên đại, nhân vật, thiếu sự kiện hoặc suy luận vượt "
        "quá tư liệu.",
        "Đã hoàn thành",
    ),
    (
        "Nút báo sai và hàng đợi phản hồi",
        "Nút “Báo câu trả lời sai” hiện thực yêu cầu “Câu trả lời này sai” và dùng mã câu trả lời "
        "do máy chủ cấp. Máy chủ lấy lại đúng "
        "câu hỏi, câu trả lời, nguồn và dữ liệu đối chiếu đã đăng ký rồi mới ghi JSONL cùng nhận xét; "
        "trạng thái ban đầu là `pending_teacher_review`. Hàng đợi không tự gắn nhãn đúng/sai.",
        "Tạo vòng phản hồi có thể truy vết mà không biến phản ánh của người dùng thành "
        "kết luận chuyên môn.",
        "Đã hoàn thành — tiếp nhận phản hồi; chưa có dashboard giáo viên",
    ),
    (
        "Thẩm định của giáo viên",
        "Xây dựng màn hình để giáo viên xem chứng cứ, đánh dấu đúng/sai/thiếu/chưa rõ, ghi "
        "lý do và duyệt báo cáo. Chỉ gắn nhãn “Đã thẩm định” sau thao tác thật của giáo viên.",
        "Tách xác nhận chuyên môn khỏi unit test kỹ thuật.",
        "Mục tiêu — chưa có dashboard và chưa có kết quả giáo viên duyệt",
    ),
    (
        "Cổng dừng khi bằng chứng dưới ngưỡng",
        "Backend dừng trước bước sinh nội dung khi kết quả truy xuất mang nhãn `below_threshold`, "
        "thay vì đoán. Ngưỡng mặc định chưa được đặt cho đến khi có ground truth để hiệu chuẩn.",
        "Tạo cơ chế từ chối đúng khi thiếu căn cứ, đồng thời không biến cosine score thành xác suất.",
        "Đã hoàn thành cơ chế; đang triển khai hiệu chuẩn ngưỡng",
    ),
    (
        "Kiểm thử mù bằng câu hỏi mới",
        "Quy trình dự kiến 20 câu chưa có trong corpus đã được lập: 12 câu đúng phạm vi nhưng "
        "diễn đạt mới, 4 câu mơ hồ và 4 câu ngoài phạm vi. Sau khi khóa corpus/câu lệnh, báo "
        "cáo riêng tỷ lệ trả lời đúng và tỷ lệ từ chối/yêu cầu làm rõ đúng.",
        "Đo khả năng khái quát và biết giới hạn; tránh đánh giá trên chính dữ liệu đã dùng "
        "để phát triển.",
        "Mục tiêu — đã có quy trình, chưa chạy và chưa có số liệu",
    ),
    (
        "Tốc độ phản hồi",
        "Giữ 23,32 giây là baseline; đo riêng truy xuất, sinh nội dung và video; tận dụng cache, "
        "song song hóa phần phù hợp, sau đó chạy lại cùng quy trình 60 câu và công bố trung bình/P95.",
        "Giảm thời gian chờ mà vẫn giữ độ đúng và khả năng truy vết.",
        "Mục tiêu — trung bình < 20 giây; chưa đạt",
    ),
    (
        "Tiến trình chờ và hủy yêu cầu",
        "Giao diện hiển thị trạng thái theo thời gian, số giây đã chờ và nút hủy. Nút hủy dùng "
        "`AbortController` để dừng yêu cầu phía client; chưa có minh chứng rằng phép tính phía server "
        "cũng bị hủy.",
        "Làm thời gian chờ dễ hiểu và trao quyền kiểm soát cho học sinh.",
        "Đã hoàn thành — trạng thái theo thời gian và hủy phía client",
    ),
    (
        "Tư liệu lịch sử xác thực — bộ thí điểm",
        "Mục “Tư liệu lịch sử xác thực” đã có 4 bản ghi cho Cách mạng tư sản Anh và Pháp. "
        "Mỗi bản ghi dùng hồ sơ Wikimedia Commons phạm vi công cộng/CC0, kèm chú thích, nơi lưu giữ, "
        "giấy phép và liên kết hồ sơ gốc.",
        "Tăng khả năng quan sát chứng cứ; tránh dùng hình minh họa không rõ xuất xứ.",
        "Đã hoàn thành — phạm vi thí điểm 2 chủ đề",
    ),
    (
        "Mở rộng kho tư liệu hình ảnh",
        "Bổ sung có kiểm duyệt bản đồ, tranh, hiện vật và văn kiện cho các chủ đề lớp 8 còn lại; "
        "nếu chưa có tư liệu đúng chủ đề và rõ quyền sử dụng thì không chèn ảnh thay thế.",
        "Duy trì tính xác thực khi mở rộng, thay vì dùng hình chỉ để trang trí.",
        "Đang triển khai",
    ),
    (
        "Lộ trình mở rộng chương trình",
        "Thí điểm một chủ đề lớp 8 được chốt trước; chỉ mở rộng toàn bộ lớp 8 sau khi đạt ngưỡng "
        "đánh giá đã thống nhất với giáo viên. Lớp 6, 7 và 9 là giai đoạn sau, không thuộc phạm vi hiện tại.",
        "Kiểm soát chất lượng theo từng giai đoạn và không tuyên bố phạm vi rộng hơn minh chứng.",
        "Mục tiêu",
    ),
    (
        "Khắc phục lượt lỗi trong benchmark",
        "Phân tích lượt lỗi HTTP 502 do Pinecone read timeout, điều chỉnh timeout/retry có giới hạn và "
        "chạy lại câu lỗi cùng toàn bộ 60 câu để kiểm tra hồi quy.",
        "Phân biệt lỗi hạ tầng với lỗi nội dung và tránh tuyên bố đã sửa khi chưa có lần đo lại.",
        "Đang triển khai — chưa xác nhận bằng benchmark mới",
    ),
]


def find_paragraph(document: Document, prefix: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Cannot find paragraph beginning with: {prefix!r}")


def replace_paragraph(paragraph, lead: str, body: str) -> None:
    paragraph.clear()
    lead_run = paragraph.add_run(lead)
    lead_run.bold = True
    paragraph.add_run(body)


def insert_paragraph_after(reference, lead: str, body: str):
    new_xml = OxmlElement("w:p")
    reference._p.addnext(new_xml)
    paragraph = reference._parent.add_paragraph()
    paragraph._p.getparent().remove(paragraph._p)
    new_xml.addnext(paragraph._p)
    new_xml.getparent().remove(new_xml)

    paragraph.style = reference.style
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.08
    replace_paragraph(paragraph, lead, body)
    return paragraph


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(width_cm))))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top: int = 70, start: int = 80, bottom: int = 70, end: int = 80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "B7A394")


def format_cell_text(cell, *, header: bool = False, status: bool = False) -> None:
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if header or status else WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(8.5 if not header else 9)
            run.bold = header or status
            if header:
                run.font.color.rgb = RGBColor(255, 255, 255)


def build_status_table(document: Document, old_table) -> None:
    table = document.add_table(rows=1, cols=4)
    table.autofit = False
    if old_table.style is not None:
        table.style = old_table.style

    widths = (3.2, 6.2, 4.3, 3.4)
    headers = ("Nội dung chỉnh sửa", "Cách cải tiến", "Ý nghĩa", "Trạng thái")
    for index, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        cell.text = text
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, "7F1D1D")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        format_cell_text(cell, header=True)
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    status_fills = {
        "Đã hoàn thành": "E2F0D9",
        "Đang triển khai": "FFF2CC",
        "Mục tiêu": "DDEBF7",
    }
    for row_index, values in enumerate(STATUS_ROWS, start=1):
        row = table.add_row()
        prevent_row_split(row)
        for cell_index, (cell, value) in enumerate(zip(row.cells, values)):
            cell.text = value
            set_cell_width(cell, widths[cell_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if cell_index == 3 else WD_CELL_VERTICAL_ALIGNMENT.TOP
            if row_index % 2 == 0 and cell_index != 3:
                set_cell_shading(cell, "F8F3EE")
            if cell_index == 3:
                for prefix, fill in status_fills.items():
                    if value.startswith(prefix):
                        set_cell_shading(cell, fill)
                        break
            format_cell_text(cell, status=cell_index == 3)

    set_table_borders(table)
    old_table._tbl.addprevious(table._tbl)
    old_table._tbl.getparent().remove(old_table._tbl)


def validate_output(path: Path) -> None:
    revised = Document(path)
    body_text = "\n".join(paragraph.text for paragraph in revised.paragraphs)
    required_text = (
        "59/60",
        "23,32 giây",
        "Mục tiêu: giảm thời gian phản hồi trung bình xuống dưới 20 giây",
        "không phải kiểm thử mù",
        "Câu trả lời này sai",
        "chưa có dashboard thẩm định",
    )
    for value in required_text:
        if value not in body_text and not any(value in cell.text for table in revised.tables for row in table.rows for cell in row.cells):
            raise AssertionError(f"Missing required text: {value}")

    expected_headers = ["Nội dung chỉnh sửa", "Cách cải tiến", "Ý nghĩa", "Trạng thái"]
    matching = [table for table in revised.tables if [cell.text for cell in table.rows[0].cells] == expected_headers]
    if len(matching) != 1:
        raise AssertionError("Expected exactly one four-column revision-status table")
    if len(matching[0].rows) != len(STATUS_ROWS) + 1:
        raise AssertionError("Revision-status table row count is incorrect")
    if len(revised.inline_shapes) != 7:
        raise AssertionError("Embedded evidence images were not preserved")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    document = Document(SOURCE)
    if len(document.tables) < 6:
        raise ValueError("Unexpected source structure: limitation table not found")

    document.core_properties.title = "Hồ sơ Bảng A HISEDU — bản chỉnh sửa theo phản biện"
    document.core_properties.subject = (
        "Phân loại rõ kết quả đã hoàn thành, nội dung đang triển khai và mục tiêu"
    )

    replace_paragraph(
        find_paragraph(document, "Trong bản thử nghiệm,"),
        "Phạm vi bản thử nghiệm hiện tại:",
        " Sản phẩm đang làm việc với kho tri thức Lịch sử lớp 8 đã nạp; chưa công bố hỗ trợ "
        "lớp 6, 7 hoặc 9. Người dùng có thể nhập câu hỏi bằng chữ hoặc tải ảnh chứa câu hỏi. Hệ "
        "thống nhận diện nội dung, tìm đoạn liên quan rồi tạo câu trả lời, sơ đồ tư duy và câu hỏi luyện "
        "tập. Việc mở rộng khối lớp chỉ thực hiện sau thử nghiệm thí điểm và thẩm định.",
    )

    replace_paragraph(
        find_paragraph(document, "Tình huống hỏi đáp:"),
        "Tình huống hỏi đáp và đối chiếu:",
        " Nhóm trình diễn bằng một câu hỏi thuộc chủ đề đã có trong kho dữ liệu. Trên màn hình, "
        "học sinh nhập câu hỏi; hệ thống hiển thị nguồn, đoạn dữ liệu dùng để đối chiếu, trạng thái thẩm "
        "định và câu trả lời ngắn theo các ý chính. Đoạn đối chiếu là nội dung trong bản ghi HISEDU, không "
        "được gọi là trích dẫn nguyên văn SGK khi trang chưa được xác minh.",
    )

    replace_paragraph(
        find_paragraph(document, "Tình huống ôn tập và kiểm tra giới hạn:"),
        "Tình huống ôn tập và kiểm tra giới hạn:",
        " Từ cùng chủ đề, học sinh có thể xem sơ đồ tư duy và làm câu hỏi luyện tập. Với Cách mạng "
        "tư sản Anh và Pháp, giao diện còn hiển thị bộ tư liệu xác thực thí điểm gồm 4 hồ sơ phạm vi công "
        "cộng/CC0 có chú thích, nơi lưu giữ và liên kết giấy phép. Một câu ngoài phạm vi cần được dùng để "
        "kiểm tra cổng dừng dưới ngưỡng; ngưỡng vẫn phải được hiệu chuẩn bằng ground truth.",
    )

    replace_paragraph(
        find_paragraph(document, "Cách kiểm tra kết quả:"),
        "Cách kiểm tra kết quả hiện có:",
        " Nhóm so sánh nội dung với đáp án tham chiếu và kiểm tra câu trả lời có đúng câu hỏi, mốc thời "
        "gian, nhân vật, diễn biến và nguồn hay không. Bộ 60 câu là kiểm thử hồi quy từ dữ liệu có cặp "
        "hỏi–đáp, không phải kiểm thử mù. Đánh giá chuyên môn của giáo viên là bước riêng và chưa thể "
        "được thay thế bằng unit test.",
    )

    replace_paragraph(
        find_paragraph(document, "Minh chứng nên lưu theo từng lượt:"),
        "Minh chứng và phản hồi sai:",
        " Mỗi kết quả hiện có mã câu trả lời do máy chủ cấp, nguồn, nội dung tham chiếu và trạng thái "
        "thẩm định. Khi nhấn “Báo câu trả lời sai”, máy chủ lưu chính câu hỏi, câu trả lời, nguồn, dữ liệu "
        "đối chiếu, nhận xét và thời điểm vào JSONL với trạng thái `pending_teacher_review`; trình duyệt "
        "không được tự gửi thay nội dung cần thẩm định. Hiện chưa có "
        "dashboard thẩm định hoàn chỉnh và chưa có báo cáo nào được gắn nhãn giáo viên đã xác minh.",
    )

    baseline = find_paragraph(document, "Số câu hỏi đã thử:")
    baseline.clear()
    baseline.add_run("Kết quả đường cơ sở (baseline, 22/09/2026): ").bold = True
    baseline.add_run(
        "60 câu được thử; 59/60 câu trả lời đạt tiêu chí (98,3%); thời gian phản hồi trung bình "
        "23,32 giây. Câu chưa đạt trả về HTTP 502 do Pinecone read timeout; đây là lỗi hạ tầng trong lần "
        "đo, không phải kết luận rằng nội dung lịch sử của câu đó sai. Các con số này được giữ nguyên làm "
        "mốc so sánh, không thay bằng số liệu mục tiêu."
    )

    replace_paragraph(
        find_paragraph(document, "Diễn giải số liệu:"),
        "Diễn giải và mục tiêu:",
        " Kết quả 59/60 chỉ phản ánh bộ kiểm thử hiện có, chưa chứng minh hệ thống ổn định với câu hỏi "
        "mới, câu ngoài phạm vi hoặc mọi loại ảnh. Mục tiêu: giảm thời gian phản hồi trung bình xuống dưới 20 giây "
        "(< 20 giây) trong cùng điều kiện đo; 23,32 giây vẫn là kết quả hiện tại cho đến khi benchmark mới "
        "được chạy và lưu minh chứng.",
    )

    section_heading = find_paragraph(document, "7. Hạn chế và hướng cải tiến")
    section_heading.text = "7. Hạn chế, phản hồi chuyên gia và kế hoạch cải tiến"
    section_heading.runs[0].bold = True
    section_heading.runs[0].font.size = Pt(13)
    insert_paragraph_after(
        section_heading,
        "Quy ước trạng thái:",
        " “Đã hoàn thành” chỉ dùng khi đã có sản phẩm hoặc minh chứng kiểm thử; “Đang triển khai” "
        "nghĩa là đang cập nhật nhưng chưa nghiệm thu; “Mục tiêu” là kết quả dự kiến, không phải thành tích hiện tại.",
    )

    build_status_table(document, document.tables[5])

    replace_paragraph(
        find_paragraph(document, "Thứ tự ưu tiên:"),
        "Thứ tự ưu tiên:",
        " (1) hiệu chuẩn cổng dừng và thực hiện quy trình kiểm thử mù 20 câu; (2) xử lý Pinecone read "
        "timeout và chạy lại benchmark 60 câu; (3) xây dashboard giáo viên trên hàng đợi phản hồi hiện có; "
        "(4) đo và tối ưu tốc độ; (5) mở rộng kho tư liệu xác thực; (6) thí điểm một chủ đề lớp 8 "
        "trước khi mở rộng.",
    )

    replace_paragraph(
        find_paragraph(document, "Kế hoạch kiểm chứng tiếp theo:"),
        "Kế hoạch kiểm thử mù:",
        " Quy trình tại reports/blind_classroom_pilot_protocol.md quy định 20 câu mới: 12 câu đúng phạm vi "
        "nhưng diễn đạt khác, 4 câu mơ hồ và 4 câu ngoài phạm vi. Khóa corpus, câu lệnh và mã nguồn "
        "trước khi nhận bộ câu hỏi từ bạn học; không "
        "cho người phát triển xem đáp án trước. Gán nhãn trước cho câu trong phạm vi và ngoài phạm vi, sau đó "
        "báo cáo riêng: số/tỷ lệ trả lời đúng, số/tỷ lệ từ chối hoặc yêu cầu làm rõ đúng, thời gian "
        "phản hồi trung bình và P95. Mọi lần sửa sau khi xem kết quả phải được ghi thành vòng kiểm thử mới.",
    )

    replace_paragraph(
        find_paragraph(document, "Kiểm tra độ tin cậy lịch sử:"),
        "Kiểm tra độ tin cậy lịch sử:",
        " AI có thể sai niên đại, nhầm nhân vật hoặc bỏ sót sự kiện ngay cả khi câu văn trôi chảy. Vì vậy, mỗi "
        "báo cáo cần cho phép đối chiếu câu trả lời với đoạn dữ liệu dùng để đối chiếu; người kiểm tra đánh dấu từng ý "
        "là chính xác, thiếu, thừa hay chưa rõ. Chỉ báo cáo là “giáo viên đã xác minh” nếu có người duyệt, "
        "thời điểm duyệt và kết luận được lưu lại.",
    )

    replace_paragraph(
        find_paragraph(document, "Trải nghiệm và dữ liệu người dùng:"),
        "Trải nghiệm, tốc độ và quyền riêng tư:",
        " Trong khi mục tiêu dưới 20 giây chưa đạt, giao diện nay hiển thị trạng thái theo thời gian, số giây "
        "đã chờ và nút hủy yêu cầu phía client. Thời gian nhận dạng ảnh, truy xuất, sinh câu trả lời và tìm video phải được "
        "đo riêng. Ảnh có tên, lớp hoặc thông tin cá nhân phải được che hoặc loại bỏ trước khi lưu phục vụ "
        "kiểm thử.",
    )

    document.save(OUTPUT)
    validate_output(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
