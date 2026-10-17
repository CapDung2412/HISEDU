"""Update the latest HISEDU evidence report without overwriting its source.

The update preserves the existing evidence images and user-authored structure,
adds the curated historical-image evidence, refreshes verified test counts, and
applies a restrained monochrome A4 layout.
"""

from __future__ import annotations

import argparse
import hashlib
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
INK = RGBColor(34, 34, 34)
MUTED = RGBColor(89, 89, 89)


def find_paragraph(document: Document, exact: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == exact:
            return paragraph
    raise ValueError(f"Không tìm thấy đoạn: {exact}")


def replace_paragraph(document: Document, prefix: str, value: str) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            paragraph.text = value
            return
    raise ValueError(f"Không tìm thấy đoạn bắt đầu bằng: {prefix}")


def new_paragraph_before(document: Document, target, text: str = "", style: str | None = None):
    paragraph = document.add_paragraph(text, style=style)
    target._p.addprevious(paragraph._p)
    return paragraph


def new_paragraph_after(document: Document, target, text: str = "", style: str | None = None):
    paragraph = document.add_paragraph(text, style=style)
    target._p.addnext(paragraph._p)
    return paragraph


def set_cell_text(cell, value: str) -> None:
    cell.text = value
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders_none(table) -> None:
    properties = table._tbl.tblPr
    existing = properties.find(qn("w:tblBorders"))
    if existing is not None:
        properties.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    properties.append(borders)


def remove_shading_and_borders(element) -> None:
    for expression in (".//w:shd", ".//w:pBdr", ".//w:tcBorders"):
        for node in element.xpath(expression):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)


def add_historical_image_section(document: Document, image_path: Path) -> None:
    target = find_paragraph(document, "5. Kết quả kiểm thử kỹ thuật")
    heading = new_paragraph_before(
        document,
        target,
        "4.6. Ảnh tư liệu đã kiểm duyệt",
        "Heading 2",
    )
    body = new_paragraph_before(
        document,
        target,
        (
            "HISEDU bổ sung kho 9 ảnh tư liệu cục bộ có hồ sơ nguồn và quyền sử dụng rõ ràng, "
            "bao phủ bảy nhóm chủ đề tiêu biểu của chương trình Lịch sử lớp 8. Bộ gợi ý chỉ trả "
            "ảnh khi câu hỏi nêu đúng sự kiện, nhân vật hoặc chủ đề tương ứng; không dùng ảnh "
            "không liên quan để lấp chỗ trống. Mỗi thẻ tư liệu hiển thị loại hình, tác giả, "
            "niên đại, đơn vị lưu giữ, giấy phép, lý do liên quan và câu hỏi gợi ý quan sát. "
            "Việc chọn ảnh chạy song song với quá trình tạo bài học và được lưu cache; ảnh được "
            "lưu cục bộ để giảm phụ thuộc máy chủ ngoài, đồng thời vẫn giữ liên kết hồ sơ gốc "
            "để đối chiếu."
        ),
    )
    bullet_1 = new_paragraph_before(
        document,
        target,
        "Giới hạn tối đa bốn ảnh cho mỗi câu hỏi; không có kết quả phù hợp thì trả danh sách rỗng.",
        "List Bullet",
    )
    bullet_2 = new_paragraph_before(
        document,
        target,
        "Đường dẫn ảnh bị giới hạn trong kho /historical/; nguồn và giấy phép được kiểm tra ngày 17/10/2026.",
        "List Bullet",
    )
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image_path), width=Cm(15.5))
    target._p.addprevious(image_paragraph._p)
    caption = new_paragraph_before(
        document,
        target,
        (
            "Hình 10. Ảnh tư liệu liên quan được hiển thị cùng nguồn, giấy phép và câu hỏi gợi ý "
            "quan sát. [Ảnh từ kho cục bộ đã kiểm duyệt; hồ sơ nguồn ghi tại "
            "web/public/historical/README.md.]"
        ),
    )
    image_paragraph.paragraph_format.keep_with_next = True
    caption.paragraph_format.keep_together = True
    for paragraph in (heading, body, bullet_1, bullet_2, caption):
        paragraph.paragraph_format.space_after = Pt(6)


def update_content(document: Document, historical_image: Path) -> None:
    # Cover metadata and safety note.
    metadata = document.tables[0]
    set_cell_text(metadata.cell(3, 1), "17/10/2026")
    scope = find_paragraph(document, "Phạm vi nội dung: Lịch sử lớp 8")
    safety = new_paragraph_after(
        document,
        scope,
        "Tài liệu không chứa khóa API, tệp .env hoặc thông tin cá nhân.",
    )
    safety.alignment = WD_ALIGN_PARAGRAPH.CENTER

    replace_paragraph(
        document,
        "HISEDU là trợ lý học tập",
        (
            "HISEDU là trợ lý học tập môn Lịch sử dành cho học sinh THCS. Hệ thống hỗ trợ đặt "
            "câu hỏi bằng văn bản hoặc hình ảnh, truy xuất dữ liệu đã lập chỉ mục, trình bày "
            "câu trả lời có nguồn, gợi ý ảnh tư liệu đã kiểm duyệt và tạo nội dung hỗ trợ ghi nhớ. "
            "Báo cáo này tập trung vào nội dung Lịch sử lớp 8."
        ),
    )
    replace_paragraph(
        document,
        "Quy trình kiểm tra ghi nhận",
        (
            "Kết quả audit ghi nhận 327/327 cặp TXT–JSON hợp lệ: toàn bộ JSON đọc được theo "
            "chế độ strict, tệp dùng UTF-8 và nội dung câu hỏi khớp giữa TXT với JSON. Không phát "
            "hiện bản ghi trùng; id, chunk_id, tên tệp và thư mục đều khớp. Có 665 đường dẫn được "
            "kiểm tra, tất cả dùng ký tự ASCII, chữ thường, không có khoảng trắng hoặc đường dẫn lỗi."
        ),
    )
    replace_paragraph(
        document,
        "Kết quả đối chiếu cho thấy 327/327",
        (
            "Manifest ingest xác nhận 327/327 bản ghi đã được upsert vào Pinecone: index "
            "hisedu-clip-v1, namespace history8-v1-0-0, độ đo cosine và vector 512 chiều. "
            "Pipeline dùng clip-ViT-B-32-multilingual-v1 cho văn bản và clip-ViT-B-32 cho ảnh; "
            "metadata đi kèm vector được giữ để truy vết nguồn khi trả lời."
        ),
    )
    replace_paragraph(
        document,
        "Phạm vi phép đo:",
        (
            "Phạm vi phép đo: 60 câu nguyên văn trong dataset, chạy hai luồng đồng thời và có "
            "tính thời gian khởi động. Kết quả không đại diện cho ảnh/OCR, câu diễn đạt mới qua "
            "CLIP/Pinecone hoặc mọi điều kiện mạng; đây chưa phải đánh giá mù độc lập."
        ),
    )
    replace_paragraph(document, "4. Năm nội dung cải tiến", "4. Sáu nội dung cải tiến")

    # Record the durable Windows dev-server correction without presenting it as
    # an application latency result.
    target_ocr = find_paragraph(document, "4.2. Nhận diện ảnh – OCR có thể chỉnh sửa")
    stability = new_paragraph_before(
        document,
        target_ocr,
        (
            "Ổn định môi trường chạy local: frontend trên Windows đã bỏ cờ Turbopack, được cấp "
            "đúng quyền Modify/Delete cho thư mục web và làm sạch cache .next. Hai lần khởi động "
            "xác nhận liên tiếp đạt trạng thái Ready sau 2,4–2,6 giây và trả HTTP 200; lỗi EPERM "
            "khi đổi tên manifest không tái diễn trong lần kiểm tra. Đây là số đo khởi động local, "
            "không phải thời gian trả lời câu hỏi của hệ thống."
        ),
    )
    stability.paragraph_format.space_after = Pt(8)

    add_historical_image_section(document, historical_image)
    replace_paragraph(
        document,
        "Hình 10. Kết quả kiểm thử",
        "Hình 11. Kết quả kiểm thử tự động của backend và frontend ngày 17/10/2026.",
    )

    # Evidence index: keep all existing evidence and split images/tests into two rows.
    evidence = document.tables[2]
    set_cell_text(evidence.cell(8, 0), "MC08")
    set_cell_text(evidence.cell(8, 1), "Ảnh tư liệu đã kiểm duyệt và hồ sơ nguồn")
    set_cell_text(evidence.cell(8, 2), "12_historical_images.png; web/public/historical/README.md")
    new_row = evidence.add_row()
    set_cell_text(new_row.cells[0], "MC09")
    set_cell_text(new_row.cells[1], "Kết quả kiểm thử và benchmark")
    set_cell_text(new_row.cells[2], "11_test_results.png; local_benchmark_60_2026-10-16_optimized.json")

    tests = document.tables[3]
    set_cell_text(tests.cell(1, 1), "87/87")
    set_cell_text(tests.cell(2, 1), "31/31")
    tests_heading = find_paragraph(document, "5. Kết quả kiểm thử kỹ thuật")
    test_note = new_paragraph_after(
        document,
        tests_heading,
        (
            "Chạy lại ngày 17/10/2026: backend đạt 87/87 kiểm thử; frontend đạt 31/31 kiểm thử "
            "trên 9 tệp; TypeScript và ESLint đều đạt. Trong đó có 12 kiểm thử backend và 4 kiểm "
            "thử frontend dành riêng cho tính năng ảnh tư liệu."
        ),
    )
    test_note.paragraph_format.space_after = Pt(6)

    limit_target = find_paragraph(
        document,
        "Mở rộng nội dung sang lớp 6, 7 và 9 sau khi hoàn tất kiểm duyệt dữ liệu lớp 8.",
    )
    image_limit = new_paragraph_before(
        document,
        limit_target,
        (
            "Kho ảnh hiện có 9 ảnh cho 7 nhóm chủ đề; câu hỏi ngoài các nhóm này sẽ không hiển "
            "thị ảnh cho đến khi tư liệu mới được bổ sung và kiểm duyệt."
        ),
        "List Bullet",
    )
    image_limit.paragraph_format.space_after = Pt(3)

    replace_paragraph(
        document,
        "Năm cải tiến giúp HISEDU",
        (
            "Sáu nội dung cải tiến giúp HISEDU phản hồi rõ ràng hơn, nhận diện ảnh an toàn hơn, "
            "lựa chọn video chính xác hơn, tổ chức kiến thức có hệ thống, hỗ trợ học sinh tự kiểm "
            "tra mức độ hiểu bài và tiếp cận ảnh tư liệu có thể đối chiếu nguồn. Các số liệu trong "
            "báo cáo gắn với tệp minh chứng và nêu rõ giới hạn để không vượt quá phạm vi kết quả "
            "thực tế."
        ),
    )


def apply_a4_monochrome_layout(document: Document) -> None:
    for section in document.sections:
        section.start_type = WD_SECTION.NEW_PAGE if section.start_type is None else section.start_type
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    # Leave enough room for the complete evidence index on page 2 so MC09 is
    # not orphaned at the top of the following page.
    if document.inline_shapes:
        cover_evidence = document.inline_shapes[0]
        scale = float(Cm(14.0)) / float(cover_evidence.width)
        cover_evidence.width = Cm(14.0)
        cover_evidence.height = int(float(cover_evidence.height) * scale)

    style_settings = {
        "Title": (22, True, 12, 12),
        "Heading 1": (16, True, 14, 8),
        "Heading 2": (13, True, 10, 6),
    }
    for style_name, (size, bold, before, after) in style_settings.items():
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = INK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        paragraph.paragraph_format.line_spacing = 1.15
        if paragraph.style.name == "Normal" and text and not text.startswith("Hình "):
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if text in {
                    "TRỢ LÝ HỌC TẬP MÔN LỊCH SỬ",
                    "HISEDU",
                    "BÁO CÁO MINH CHỨNG SẢN PHẨM",
                    "Phạm vi nội dung: Lịch sử lớp 8",
                    "Tài liệu không chứa khóa API, tệp .env hoặc thông tin cá nhân.",
                }
                else WD_ALIGN_PARAGRAPH.JUSTIFY
            )
            paragraph.paragraph_format.space_after = Pt(6)
        elif text.startswith("Hình "):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.keep_together = True
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.color.rgb = MUTED if text.startswith("Hình ") else INK
            run.font.underline = False

    for table_index, table in enumerate(document.tables):
        table.style = "Normal Table"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_borders_none(table)
        remove_shading_and_borders(table._tbl)
        # Re-add explicit nil borders after removing all legacy border nodes.
        set_table_borders_none(table)
        for row_index, row in enumerate(table.rows):
            row_properties = row._tr.get_or_add_trPr()
            if row_properties.find(qn("w:cantSplit")) is None:
                row_properties.append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run._element.get_or_add_rPr().rFonts.set(
                            qn("w:eastAsia"), "Times New Roman"
                        )
                        run.font.size = Pt(9.5 if table_index == 2 else 10.5)
                        run.font.color.rgb = INK
                        run.font.bold = row_index == 0
                        run.font.underline = False


def strip_decorative_xml(xml_bytes: bytes) -> bytes:
    root = etree.fromstring(xml_bytes)
    for expression in (".//w:shd", ".//w:pBdr", ".//w:tcBorders"):
        for node in root.xpath(expression, namespaces=NS):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
    for node in root.xpath(".//w:color", namespaces=NS):
        node.set(f"{{{W_NS}}}val", "222222")
        for attribute in ("themeColor", "themeTint", "themeShade"):
            node.attrib.pop(f"{{{W_NS}}}{attribute}", None)
    for node in root.xpath(".//w:u", namespaces=NS):
        node.set(f"{{{W_NS}}}val", "none")
    for table in root.xpath(".//w:tbl", namespaces=NS):
        properties = table.find(f"{{{W_NS}}}tblPr")
        if properties is None:
            properties = etree.Element(f"{{{W_NS}}}tblPr")
            table.insert(0, properties)
        existing = properties.find(f"{{{W_NS}}}tblBorders")
        if existing is not None:
            properties.remove(existing)
        borders = etree.SubElement(properties, f"{{{W_NS}}}tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            item = etree.SubElement(borders, f"{{{W_NS}}}{edge}")
            item.set(f"{{{W_NS}}}val", "nil")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def replace_test_image_and_clean_xml(source: Path, output: Path, test_image: Path) -> None:
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with zipfile.ZipFile(source, "r") as source_zip, zipfile.ZipFile(temporary, "w") as target_zip:
            target_zip.comment = source_zip.comment
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == "word/media/image10.png":
                    data = test_image.read_bytes()
                elif info.filename == "word/document.xml":
                    data = strip_decorative_xml(data)
                target_zip.writestr(info, data)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)


def validate_report(path: Path, historical_image: Path, test_image: Path) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        if archive.testzip() is not None:
            raise ValueError("Gói DOCX bị lỗi CRC")
        xml = archive.read("word/document.xml").decode("utf-8")
        required = (
            "327/327",
            "60/60",
            "10,167 giây",
            "87/87",
            "31/31",
            "4.6. Ảnh tư liệu đã kiểm duyệt",
            "9 ảnh tư liệu",
            "bảy nhóm chủ đề",
            "Sáu nội dung cải tiến",
            "Tài liệu không chứa khóa API",
        )
        for value in required:
            if value not in xml:
                raise ValueError(f"Báo cáo thiếu nội dung bắt buộc: {value}")
        for stale in ("75/75", "29/29", "4. Năm nội dung cải tiến", "Phạm vi phép đo:60"):
            if stale in xml:
                raise ValueError(f"Báo cáo còn nội dung cũ: {stale}")
        if "w:shd" in xml or "w:pBdr" in xml or "w:tcBorders" in xml:
            raise ValueError("Báo cáo còn nền hoặc đường kẻ trang trí")
        embedded_test = archive.read("word/media/image10.png")
        if hashlib.sha256(embedded_test).digest() != hashlib.sha256(test_image.read_bytes()).digest():
            raise ValueError("Ảnh kiểm thử chưa được cập nhật")
        media_hashes = {
            hashlib.sha256(archive.read(name)).digest()
            for name in archive.namelist()
            if name.startswith("word/media/")
        }
        if hashlib.sha256(historical_image.read_bytes()).digest() not in media_hashes:
            raise ValueError("Ảnh minh chứng tư liệu chưa được nhúng")

    document = Document(path)
    if len(document.tables) != 4 or len(document.inline_shapes) != 11:
        raise ValueError(
            f"Cấu trúc không đúng: {len(document.tables)} bảng, {len(document.inline_shapes)} ảnh"
        )
    section = document.sections[0]
    if round(section.page_width.cm, 1) != 21.0 or round(section.page_height.cm, 1) != 29.7:
        raise ValueError("Báo cáo chưa dùng khổ A4")
    if len(document.tables[2].rows) != 10:
        raise ValueError("Danh mục minh chứng chưa có MC09")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Cập nhật báo cáo minh chứng HISEDU 17/10/2026")
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--historical-image", type=Path, required=True)
    parser.add_argument("--test-image", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    source = args.source.resolve()
    historical_image = args.historical_image.resolve()
    test_image = args.test_image.resolve()
    output = args.output.resolve()
    for path in (source, historical_image, test_image):
        if not path.is_file():
            raise FileNotFoundError(path)
    output.parent.mkdir(parents=True, exist_ok=True)

    document = Document(source)
    update_content(document, historical_image)
    apply_a4_monochrome_layout(document)
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".docx", delete=False) as handle:
        staged = Path(handle.name)
    try:
        document.save(staged)
        replace_test_image_and_clean_xml(staged, output, test_image)
    finally:
        staged.unlink(missing_ok=True)
    validate_report(output, historical_image, test_image)
    print(f"UPDATED_REPORT={output}")


if __name__ == "__main__":
    main()
